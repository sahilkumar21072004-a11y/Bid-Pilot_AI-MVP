import json
import logging
import re
import io
import hmac
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response
from pydantic import BaseModel, Field

from .config import settings
from .retrieval import retriever
from .security import hash_password, issue_token, verify_password, verify_token
from .workflow import run_proposal_workflow
from .store import connect, initialize, row_dict

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s %(message)s")
logger = logging.getLogger("bidpilot")
@asynccontextmanager
async def lifespan(_: FastAPI):
    initialize()
    yield


app = FastAPI(title=settings.app_name, version="0.1.0", description="Local-first RFP proposal automation MVP", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=[settings.frontend_origin, "http://127.0.0.1:5173"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])


@app.middleware("http")
async def require_workspace_auth(request: Request, call_next):
    public_paths = {"/api/health", "/api/auth/config", "/api/auth/login"}
    if not settings.auth_enabled or not request.url.path.startswith("/api/") or request.url.path in public_paths or request.method == "OPTIONS":
        return await call_next(request)
    header = request.headers.get("Authorization", "")
    user = verify_token(header[7:]) if header.startswith("Bearer ") else None
    if not user:
        return JSONResponse({"detail": "Sign in required"}, status_code=401)
    if user["role"] != "admin":
        with connect() as db:
            account = db.execute("SELECT role,active FROM team WHERE username=?", (user["username"],)).fetchone()
        if not account or not account["active"] or str(account["role"]).lower() != user["role"]:
            return JSONResponse({"detail": "Account is inactive or its permissions changed. Sign in again."}, status_code=401)
    read_export = request.method == "POST" and request.url.path.endswith("/export")
    if user["role"] == "viewer" and request.method not in {"GET", "HEAD"} and not read_export:
        return JSONResponse({"detail": "Viewer access is read-only"}, status_code=403)
    if request.url.path.startswith("/api/team") and user["role"] != "admin":
        return JSONResponse({"detail": "Administrator access required"}, status_code=403)
    if request.url.path == "/api/audit" and user["role"] != "admin":
        return JSONResponse({"detail": "Administrator access required"}, status_code=403)
    if user["role"] == "reviewer" and request.method not in {"GET", "HEAD"}:
        if not request.url.path.startswith("/api/proposals") or request.method == "DELETE":
            return JSONResponse({"detail": "Reviewers can only create or review proposals"}, status_code=403)
        match = re.match(r"^/api/proposals/(\d+)(?:/|$)", request.url.path)
        if match:
            with connect() as db:
                proposal = db.execute("SELECT reviewer FROM proposals WHERE id=?", (int(match.group(1)),)).fetchone()
                account = db.execute("SELECT name FROM team WHERE username=?", (user["username"],)).fetchone()
            assigned = proposal["reviewer"] if proposal else ""
            if assigned and assigned not in {user["username"], account["name"] if account else ""}:
                return JSONResponse({"detail": "This proposal is assigned to another reviewer"}, status_code=403)
    request.state.user = user
    return await call_next(request)


class ProposalUpdate(BaseModel):
    title: str | None = None
    sections: list[dict[str, Any]] | None = None
    requirements: list[dict[str, Any]] | None = None
    status: str | None = None
    reviewer: str | None = None
    comments: list[dict[str, str]] | None = None


class KnowledgeInput(BaseModel):
    title: str = Field(min_length=1, max_length=160)
    content: str = Field(min_length=1, max_length=20000)
    category: str = "General"
    tags: str = ""


class LoginInput(BaseModel):
    username: str = Field(min_length=1, max_length=100)
    password: str = Field(min_length=1, max_length=500)


class TeamInput(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    email: str = Field(min_length=3, max_length=254)
    username: str = Field(min_length=3, max_length=80, pattern=r"^[a-zA-Z0-9_.-]+$")
    password: str = Field(min_length=12, max_length=500)
    role: str = "Reviewer"


class TeamUpdate(BaseModel):
    role: str | None = None
    active: bool | None = None
    password: str | None = Field(default=None, min_length=12, max_length=500)


def actor_name(request: Request) -> str:
    user = getattr(request.state, "user", None)
    return user["username"] if user else "local-demo"


def write_audit(actor: str, action: str, entity_type: str, entity_id: str, details: dict[str, Any] | None = None) -> None:
    with connect() as db:
        db.execute("INSERT INTO audit_events(actor,action,entity_type,entity_id,details) VALUES(?,?,?,?,?)", (actor, action, entity_type, entity_id, json.dumps(details or {})))


def version_snapshot(proposal: dict[str, Any]) -> str:
    fields = ("title", "source_name", "requirements", "sections", "status", "reviewer", "comments", "workflow")
    return json.dumps({key: proposal[key] for key in fields}, ensure_ascii=False)


def get_proposal(proposal_id: int) -> dict[str, Any]:
    with connect() as db:
        row = db.execute("SELECT * FROM proposals WHERE id=?", (proposal_id,)).fetchone()
    item = row_dict(row)
    if not item:
        raise HTTPException(404, "Proposal not found")
    for key in ("requirements", "sections", "comments", "workflow"):
        item[key] = json.loads(item[key])
    return item


async def read_upload(file: UploadFile) -> str:
    name = Path(file.filename or "upload").name
    content = await file.read()
    if len(content) > 25 * 1024 * 1024:
        raise HTTPException(413, "File must be 25 MB or smaller")
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".pdf":
            with fitz.open(stream=content, filetype="pdf") as pdf:
                extracted = "\n".join(page.get_text() for page in pdf).strip()
                if len(extracted) >= 80:
                    return extracted
                try:
                    import pytesseract
                    from PIL import Image
                    if settings.tesseract_cmd:
                        pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd
                    ocr_pages = []
                    for page in pdf:
                        pixmap = page.get_pixmap(dpi=220, alpha=False)
                        image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                        ocr_pages.append(pytesseract.image_to_string(image))
                    ocr_text = "\n".join(ocr_pages).strip()
                    if ocr_text:
                        logger.info("OCR extracted text from scanned PDF %s", name)
                        return ocr_text
                except Exception as exc:
                    logger.info("Local OCR unavailable or unsuccessful for %s: %s", name, exc)
                if extracted:
                    return extracted
                raise HTTPException(422, "No selectable text found. Install Tesseract OCR and set TESSERACT_CMD to process this scanned PDF.")
        if suffix == ".docx":
            import io
            document = Document(io.BytesIO(content))
            return "\n".join(p.text for p in document.paragraphs)
        if suffix in {".txt", ".md"}:
            return content.decode("utf-8", errors="replace")
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Document parsing failed for %s", name)
        raise HTTPException(422, "Unable to parse document") from exc
    raise HTTPException(415, "Supported files: PDF, DOCX, TXT, MD")


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"application": settings.app_name, "status": "healthy"}


@app.get("/api/auth/config")
def auth_config(request: Request) -> dict[str, Any]:
    header = request.headers.get("Authorization", "")
    user = verify_token(header[7:]) if header.startswith("Bearer ") else None
    return {"required": settings.auth_enabled, "authenticated": bool(user) if settings.auth_enabled else True, "username": user["username"] if user else "Sahil Kumar", "role": user["role"] if user else ("admin" if not settings.auth_enabled else "")}


@app.post("/api/auth/login")
def auth_login(payload: LoginInput) -> dict[str, str]:
    if not settings.auth_enabled:
        raise HTTPException(400, "Login is disabled. Set AUTH_ENABLED=true to enable it.")
    if not settings.admin_password or not settings.session_secret:
        raise HTTPException(503, "Set ADMIN_PASSWORD and SESSION_SECRET in backend/.env before enabling login.")
    role = "admin" if hmac.compare_digest(payload.username, settings.admin_username) and hmac.compare_digest(payload.password, settings.admin_password) else ""
    if not role:
        with connect() as db:
            row = db.execute("SELECT username,password_hash,role,active FROM team WHERE username=?", (payload.username,)).fetchone()
        if row and row["active"] and row["password_hash"] and verify_password(payload.password, row["password_hash"]):
            role = str(row["role"]).lower()
    if not role:
        raise HTTPException(401, "Incorrect username or password")
    return {"access_token": issue_token(payload.username, role), "token_type": "bearer", "username": payload.username, "role": role}


@app.get("/api/dashboard")
def dashboard() -> dict[str, Any]:
    with connect() as db:
        proposals = db.execute("SELECT COUNT(*) FROM proposals").fetchone()[0]
        knowledge = db.execute("SELECT COUNT(*) FROM knowledge").fetchone()[0]
        rows = db.execute("SELECT requirements FROM proposals").fetchall()
    requirements = [item for row in rows for item in json.loads(row[0])]
    mandatory = sum(item.get("category") == "Mandatory" for item in requirements)
    reviewed = sum(item.get("review_status") == "Approved" for item in requirements)
    return {"proposals": proposals, "knowledge_entries": knowledge, "requirements": len(requirements), "mandatory": mandatory, "compliance": round(100 * reviewed / len(requirements)) if requirements else 0, "agents": [{"name": name, "status": "Ready"} for name in ["Extraction Agent", "Analysis Agent", "Drafting Agent", "Cost Estimation Agent", "Compliance Agent", "Optimisation Agent"]]}


@app.post("/api/proposals/generate")
async def generate_proposal(request: Request, file: UploadFile = File(...), template_id: int | None = None) -> dict[str, Any]:
    text = await read_upload(file)
    with connect() as db:
        company_row = db.execute("SELECT value FROM settings WHERE key='company'").fetchone()
        company = json.loads(company_row[0]) if company_row else {}
        knowledge_entries = [dict(row) for row in db.execute("SELECT * FROM knowledge ORDER BY id DESC").fetchall()]
        knowledge = retriever.retrieve(text, knowledge_entries, limit=5)
        template_row = db.execute("SELECT * FROM templates WHERE id=?", (template_id,)).fetchone() if template_id else None
        template = dict(template_row) if template_row else None
        if template:
            template["sections"] = json.loads(template["sections"])
        title = Path(file.filename or "RFP").stem
    workflow = await run_proposal_workflow(text, title, company, knowledge, template)
    requirements = workflow.get("requirements", [])
    sections = workflow.get("sections", [])
    with connect() as db:
        cursor = db.execute("INSERT INTO proposals(title,source_name,requirements,sections,workflow) VALUES(?,?,?,?,?) RETURNING id", (title, file.filename or "RFP", json.dumps(requirements), json.dumps(sections), json.dumps({"agents": workflow.get("agents", []), "llm_mode": workflow.get("llm_mode", "mock"), "compliance": workflow.get("compliance", {}), "cost_estimate": workflow.get("cost_estimate", {})})))
        proposal_id = cursor.lastrowid
    logger.info("Generated local demo proposal id=%s", proposal_id)
    result = get_proposal(int(proposal_id))
    with connect() as db:
        db.execute("INSERT INTO proposal_versions(proposal_id,version_number,snapshot,created_by) VALUES(?,?,?,?)", (proposal_id, 1, version_snapshot(result), actor_name(request)))
    write_audit(actor_name(request), "proposal.generated", "proposal", str(proposal_id), {"source_name": result["source_name"]})
    return result


@app.get("/api/proposals")
def list_proposals() -> list[dict[str, Any]]:
    with connect() as db:
        ids = [row[0] for row in db.execute("SELECT id FROM proposals ORDER BY updated_at DESC").fetchall()]
    return [get_proposal(proposal_id) for proposal_id in ids]


@app.get("/api/proposals/{proposal_id}")
def proposal_detail(proposal_id: int) -> dict[str, Any]:
    return get_proposal(proposal_id)


@app.put("/api/proposals/{proposal_id}")
def update_proposal(proposal_id: int, payload: ProposalUpdate, request: Request) -> dict[str, Any]:
    current = get_proposal(proposal_id)
    for key in ("title", "sections", "requirements", "status", "reviewer", "comments"):
        value = getattr(payload, key)
        if value is not None:
            current[key] = value
    with connect() as db:
        db.execute("UPDATE proposals SET title=?,requirements=?,sections=?,status=?,reviewer=?,comments=?,updated_at=CURRENT_TIMESTAMP WHERE id=?", (current["title"], json.dumps(current["requirements"]), json.dumps(current["sections"]), current["status"], current["reviewer"], json.dumps(current["comments"]), proposal_id))
        version_number = db.execute("SELECT COALESCE(MAX(version_number),0)+1 FROM proposal_versions WHERE proposal_id=?", (proposal_id,)).fetchone()[0]
        db.execute("INSERT INTO proposal_versions(proposal_id,version_number,snapshot,created_by) VALUES(?,?,?,?)", (proposal_id, version_number, version_snapshot(current), actor_name(request)))
        db.execute("INSERT INTO audit_events(actor,action,entity_type,entity_id,details) VALUES(?,?,?,?,?)", (actor_name(request), "proposal.updated", "proposal", str(proposal_id), json.dumps({"version": version_number})))
    return get_proposal(proposal_id)


@app.delete("/api/proposals/{proposal_id}")
def delete_proposal(proposal_id: int, request: Request) -> dict[str, bool]:
    with connect() as db:
        db.execute("INSERT INTO audit_events(actor,action,entity_type,entity_id,details) VALUES(?,?,?,?,?)", (actor_name(request), "proposal.deleted", "proposal", str(proposal_id), "{}"))
        cursor = db.execute("DELETE FROM proposals WHERE id=?", (proposal_id,))
    if not cursor.rowcount:
        raise HTTPException(404, "Proposal not found")
    return {"deleted": True}


@app.get("/api/proposals/{proposal_id}/versions")
def proposal_versions(proposal_id: int) -> list[dict[str, Any]]:
    get_proposal(proposal_id)
    with connect() as db:
        rows = db.execute("SELECT id,version_number,created_by,created_at FROM proposal_versions WHERE proposal_id=? ORDER BY version_number DESC", (proposal_id,)).fetchall()
    return [dict(row) for row in rows]


@app.post("/api/proposals/{proposal_id}/versions/{version_id}/restore")
def restore_proposal_version(proposal_id: int, version_id: int, request: Request) -> dict[str, Any]:
    get_proposal(proposal_id)
    with connect() as db:
        row = db.execute("SELECT snapshot FROM proposal_versions WHERE id=? AND proposal_id=?", (version_id, proposal_id)).fetchone()
        if not row:
            raise HTTPException(404, "Proposal version not found")
        snapshot = json.loads(row["snapshot"])
        db.execute("UPDATE proposals SET title=?,requirements=?,sections=?,status=?,reviewer=?,comments=?,workflow=?,updated_at=CURRENT_TIMESTAMP WHERE id=?", (snapshot["title"], json.dumps(snapshot["requirements"]), json.dumps(snapshot["sections"]), snapshot["status"], snapshot["reviewer"], json.dumps(snapshot["comments"]), json.dumps(snapshot["workflow"]), proposal_id))
        version_number = db.execute("SELECT COALESCE(MAX(version_number),0)+1 FROM proposal_versions WHERE proposal_id=?", (proposal_id,)).fetchone()[0]
        db.execute("INSERT INTO proposal_versions(proposal_id,version_number,snapshot,created_by) VALUES(?,?,?,?)", (proposal_id, version_number, json.dumps(snapshot), actor_name(request)))
        db.execute("INSERT INTO audit_events(actor,action,entity_type,entity_id,details) VALUES(?,?,?,?,?)", (actor_name(request), "proposal.version_restored", "proposal", str(proposal_id), json.dumps({"source_version_id": version_id, "new_version": version_number})))
    return get_proposal(proposal_id)


@app.get("/api/knowledge")
def list_knowledge(q: str = "") -> list[dict[str, Any]]:
    with connect() as db:
        rows = db.execute("SELECT * FROM knowledge WHERE title LIKE ? OR content LIKE ? OR tags LIKE ? ORDER BY id DESC", tuple([f"%{q}%"] * 3)).fetchall()
    return [dict(row) for row in rows]


@app.post("/api/knowledge")
def add_knowledge(payload: KnowledgeInput) -> dict[str, Any]:
    with connect() as db:
        cursor = db.execute("INSERT INTO knowledge(title,content,category,tags) VALUES(?,?,?,?) RETURNING id", (payload.title, payload.content, payload.category, payload.tags))
        row = db.execute("SELECT * FROM knowledge WHERE id=?", (cursor.lastrowid,)).fetchone()
    return dict(row)


@app.put("/api/knowledge/{entry_id}")
def update_knowledge(entry_id: int, payload: KnowledgeInput) -> dict[str, Any]:
    with connect() as db:
        cursor = db.execute("UPDATE knowledge SET title=?,content=?,category=?,tags=? WHERE id=?", (payload.title, payload.content, payload.category, payload.tags, entry_id))
        row = db.execute("SELECT * FROM knowledge WHERE id=?", (entry_id,)).fetchone()
    if not cursor.rowcount or not row:
        raise HTTPException(404, "Knowledge entry not found")
    return dict(row)


@app.delete("/api/knowledge/{entry_id}")
def delete_knowledge(entry_id: int) -> dict[str, bool]:
    with connect() as db:
        cursor = db.execute("DELETE FROM knowledge WHERE id=?", (entry_id,))
    if not cursor.rowcount:
        raise HTTPException(404, "Knowledge entry not found")
    return {"deleted": True}


@app.get("/api/company-profile")
def get_company() -> dict[str, Any]:
    with connect() as db:
        row = db.execute("SELECT value FROM settings WHERE key='company'").fetchone()
    return json.loads(row[0]) if row else {}


@app.put("/api/company-profile")
def save_company(payload: dict[str, Any]) -> dict[str, Any]:
    with connect() as db:
        db.execute("INSERT INTO settings(key,value) VALUES('company',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (json.dumps(payload),))
    return payload


@app.get("/api/templates")
def list_templates() -> list[dict[str, Any]]:
    with connect() as db:
        rows = db.execute("SELECT * FROM templates ORDER BY id").fetchall()
    return [{**dict(row), "sections": json.loads(row["sections"])} for row in rows]


@app.post("/api/templates")
def add_template(payload: dict[str, Any]) -> dict[str, Any]:
    name = str(payload.get("name", "")).strip()
    sections = payload.get("sections", [])
    if not name or not isinstance(sections, list) or not sections:
        raise HTTPException(422, "Template name and at least one section are required")
    with connect() as db:
        cursor = db.execute("INSERT INTO templates(name,sections) VALUES(?,?) RETURNING id", (name, json.dumps(sections)))
        row = db.execute("SELECT * FROM templates WHERE id=?", (cursor.lastrowid,)).fetchone()
    return {**dict(row), "sections": sections}


@app.delete("/api/templates/{template_id}")
def delete_template(template_id: int) -> dict[str, bool]:
    with connect() as db:
        cursor = db.execute("DELETE FROM templates WHERE id=?", (template_id,))
    if not cursor.rowcount:
        raise HTTPException(404, "Template not found")
    return {"deleted": True}


@app.get("/api/team")
def list_team() -> list[dict[str, Any]]:
    with connect() as db:
        return [dict(row) for row in db.execute("SELECT id,name,email,username,role,active FROM team ORDER BY name").fetchall()]


@app.post("/api/team")
def add_team(payload: TeamInput, request: Request) -> dict[str, Any]:
    name, email, username, role = payload.name.strip(), payload.email.strip().lower(), payload.username.strip(), payload.role
    if role not in {"Reviewer", "Viewer"}:
        raise HTTPException(422, "Team accounts may use Reviewer or Viewer role; the workspace admin is configured in environment settings")
    try:
        with connect() as db:
            cursor = db.execute("INSERT INTO team(name,email,role,username,password_hash,active) VALUES(?,?,?,?,?,?) RETURNING id", (name, email, role, username, hash_password(payload.password), True))
            member_id = cursor.lastrowid
            row = db.execute("SELECT id,name,email,username,role,active FROM team WHERE id=?", (member_id,)).fetchone()
            db.execute("INSERT INTO audit_events(actor,action,entity_type,entity_id,details) VALUES(?,?,?,?,?)", (actor_name(request), "team.member_created", "team_member", str(member_id), json.dumps({"username": username, "role": role})))
        return dict(row)
    except Exception as exc:
        logger.exception("Could not create team member")
        raise HTTPException(409, "That username or email is already in the team") from exc


@app.patch("/api/team/{member_id}")
def update_team_member(member_id: int, payload: TeamUpdate, request: Request) -> dict[str, Any]:
    if payload.role is not None and payload.role not in {"Reviewer", "Viewer"}:
        raise HTTPException(422, "Role must be Reviewer or Viewer")
    with connect() as db:
        current = db.execute("SELECT id,name,email,username,role,active FROM team WHERE id=?", (member_id,)).fetchone()
        if not current:
            raise HTTPException(404, "Team member not found")
        role = payload.role if payload.role is not None else current["role"]
        active = payload.active if payload.active is not None else current["active"]
        if payload.password:
            db.execute("UPDATE team SET role=?,active=?,password_hash=? WHERE id=?", (role, active, hash_password(payload.password), member_id))
        else:
            db.execute("UPDATE team SET role=?,active=? WHERE id=?", (role, active, member_id))
        db.execute("INSERT INTO audit_events(actor,action,entity_type,entity_id,details) VALUES(?,?,?,?,?)", (actor_name(request), "team.member_updated", "team_member", str(member_id), json.dumps({"role": role, "active": active})))
        row = db.execute("SELECT id,name,email,username,role,active FROM team WHERE id=?", (member_id,)).fetchone()
    return dict(row)


@app.delete("/api/team/{member_id}")
def remove_team(member_id: int, request: Request) -> dict[str, bool]:
    with connect() as db:
        db.execute("INSERT INTO audit_events(actor,action,entity_type,entity_id,details) VALUES(?,?,?,?,?)", (actor_name(request), "team.member_deleted", "team_member", str(member_id), "{}"))
        cursor = db.execute("DELETE FROM team WHERE id=?", (member_id,))
    if not cursor.rowcount:
        raise HTTPException(404, "Team member not found")
    return {"deleted": True}


@app.get("/api/audit")
def list_audit_events(limit: int = 100) -> list[dict[str, Any]]:
    limit = min(max(limit, 1), 500)
    with connect() as db:
        rows = db.execute("SELECT id,actor,action,entity_type,entity_id,details,created_at FROM audit_events ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
    return [{**dict(row), "details": json.loads(row["details"])} for row in rows]


@app.post("/api/proposals/{proposal_id}/export")
def export_proposal(proposal_id: int, format: str = "docx") -> Response:
    proposal = get_proposal(proposal_id)
    if format not in {"docx", "pdf"}:
        raise HTTPException(422, "format must be docx or pdf")
    import io
    if format == "docx":
        document = Document()
        document.add_heading(proposal["title"], 0)
        for section in proposal["sections"]:
            document.add_heading(section["heading"], level=1)
            document.add_paragraph(section["body"])
        document.add_heading("Requirements and compliance", level=1)
        for requirement in proposal["requirements"]:
            document.add_paragraph(f"{requirement['category']} — {requirement['text']} — {requirement['review_status']}", style="List Bullet")
        output = io.BytesIO(); document.save(output)
        return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f"attachment; filename=proposal-{proposal_id}.docx"})
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    output = io.BytesIO(); pdf = canvas.Canvas(output, pagesize=letter); y = 750
    pdf.setFont("Helvetica-Bold", 18); pdf.drawString(48, y, proposal["title"][:70]); y -= 40
    for section in proposal["sections"]:
        if y < 80: pdf.showPage(); y = 750
        pdf.setFont("Helvetica-Bold", 13); pdf.drawString(48, y, section["heading"][:80]); y -= 22
        pdf.setFont("Helvetica", 9)
        for line in re.findall(r".{1,100}(?:\s|$)", section["body"]):
            if y < 60: pdf.showPage(); y = 750; pdf.setFont("Helvetica", 9)
            pdf.drawString(48, y, line.strip()[:110]); y -= 14
        y -= 12
    pdf.save()
    return Response(output.getvalue(), media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename=proposal-{proposal_id}.pdf"})
